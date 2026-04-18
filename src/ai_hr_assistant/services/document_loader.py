from __future__ import annotations

import hashlib
import re
from pathlib import Path

from ai_hr_assistant.domain.models import Document

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


def load_documents(documents_dir: Path) -> list[Document]:
    documents: list[Document] = []

    for file_path in sorted(documents_dir.glob("*")):
        if not file_path.is_file() or file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        text = extract_text(file_path).strip()
        if not text:
            continue

        checksum = hashlib.sha256(text.encode("utf-8")).hexdigest()
        document_id = hashlib.sha1(str(file_path).encode("utf-8")).hexdigest()
        documents.append(
            Document(
                document_id=document_id,
                title=file_path.name,
                source_path=str(file_path),
                text=text,
                checksum=checksum,
            )
        )

    return documents


def extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()

    if suffix in {".txt", ".md"}:
        return file_path.read_text(encoding="utf-8")
    if suffix == ".pdf":
        return _extract_pdf_text(file_path)
    if suffix == ".docx":
        return _extract_docx_text(file_path)

    return ""


def split_text(text: str, chunk_size: int = 1200, overlap: int = 150) -> list[str]:
    normalized_text = re.sub(r"\r\n?", "\n", text)
    paragraphs = [item.strip() for item in re.split(r"\n{2,}", normalized_text) if item.strip()]
    chunks: list[str] = []
    current_chunk = ""

    for paragraph in paragraphs:
        paragraph = re.sub(r"\n+", "\n", paragraph)

        if len(paragraph) <= chunk_size:
            candidate = f"{current_chunk}\n\n{paragraph}".strip() if current_chunk else paragraph
            if len(candidate) <= chunk_size:
                current_chunk = candidate
                continue

        if current_chunk:
            chunks.append(current_chunk.strip())
            current_chunk = current_chunk[-overlap:].strip()

        if len(paragraph) <= chunk_size:
            current_chunk = f"{current_chunk}\n\n{paragraph}".strip() if current_chunk else paragraph
            continue

        start = 0
        while start < len(paragraph):
            end = start + chunk_size
            chunk = paragraph[start:end].strip()
            if chunk:
                chunks.append(chunk)
            if end >= len(paragraph):
                break
            start = max(end - overlap, start + 1)

        current_chunk = ""

    if current_chunk:
        chunks.append(current_chunk.strip())

    return chunks


def _extract_pdf_text(file_path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise RuntimeError("Для чтения PDF установите зависимость pypdf.") from error

    reader = PdfReader(str(file_path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(page.strip() for page in pages if page.strip())


def _extract_docx_text(file_path: Path) -> str:
    try:
        from docx import Document as DocxDocument
    except ImportError as error:
        raise RuntimeError("Для чтения DOCX установите зависимость python-docx.") from error

    document = DocxDocument(str(file_path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return "\n\n".join(paragraphs)
