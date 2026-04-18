from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT_DIR = Path(__file__).resolve().parents[1]
DOCUMENTS_DIR = ROOT_DIR / "documents"


def _configure_styles(document: Document) -> None:
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal_style.font.size = Pt(11)

    title_style = document.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def build_doc(
    filename: str,
    title: str,
    subtitle: str,
    purpose: str,
    sections: list[dict[str, object]],
) -> Path:
    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)

    document = Document()
    _configure_styles(document)

    title_paragraph = document.add_paragraph()
    title_paragraph.style = "Title"
    title_paragraph.alignment = 1
    title_paragraph.add_run(title)

    subtitle_paragraph = document.add_paragraph()
    subtitle_paragraph.alignment = 1
    subtitle_paragraph.add_run(subtitle)

    document.add_paragraph("")
    document.add_paragraph(
        "Демонстрационный внутренний документ для MVP AI HR Assistant. "
        "Содержание приближено к практике российских компаний среднего размера."
    )
    document.add_paragraph(f"Назначение документа: {purpose}")
    document.add_paragraph(
        "Документ предназначен для внутреннего использования сотрудниками, "
        "руководителями и HR-подразделением."
    )
    document.add_page_break()

    for index, section in enumerate(sections):
        document.add_heading(str(section["heading"]), level=1)

        for paragraph in section.get("paragraphs", []):
            document.add_paragraph(str(paragraph))

        for bullet in section.get("bullets", []):
            document.add_paragraph(str(bullet), style="List Bullet")

        if index < len(sections) - 1:
            page_break = document.add_paragraph()
            page_break.add_run().add_break(WD_BREAK.PAGE)

    target_path = DOCUMENTS_DIR / filename
    document.save(target_path)
    return target_path
